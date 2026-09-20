"""Builds small real PDF/DOCX files in memory for extraction tests."""

import io
import zipfile

import docx
from pypdf import PdfReader, PdfWriter


def make_pdf(lines: list[str]) -> bytes:
    """A minimal valid one-page PDF with the given lines of Helvetica text."""

    def esc(text: str) -> str:
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    ops = ["BT", "/F1 11 Tf", "14 TL", "50 780 Td"]
    for line in lines:
        ops += [f"({esc(line)}) Tj", "T*"]
    ops.append("ET")
    stream = "\n".join(ops).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def make_encrypted_pdf(lines: list[str], password: str = "secret") -> bytes:
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(make_pdf(lines))))
    writer.encrypt(user_password=password, owner_password=password)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def make_docx(
    paragraphs: list[str], header: str | None = None, table: list[list[str]] | None = None
) -> bytes:
    document = docx.Document()
    if header is not None:
        document.sections[0].header.paragraphs[0].text = header
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def make_zip_bomb_docx(uncompressed_mb: int = 40) -> bytes:
    """A 'DOCX' whose entries expand far beyond the extraction limit."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", b"\0" * (uncompressed_mb * 1024 * 1024))
    return out.getvalue()


CV_LINES = [
    "Frontend Developer",
    "Skills: JavaScript, TypeScript, React, Next.js, Tailwind CSS, Git",
    "Experience: Frontend intern at Paystack-like fintech, Jan 2024 to present.",
    "Built dashboard components in React and wrote unit tests with Jest.",
    "Projects: Market price tracker for Lagos traders, built with Next.js and PostgreSQL.",
    "Education: BSc Computer Science, University of Lagos, 2023.",
    "Volunteer: taught HTML and CSS basics to secondary school students every weekend.",
]
