"""Text extraction from untrusted CV files, with limits (ADR-0010).

Only body text is read. DOCX headers and footers are skipped on purpose: they usually hold contact
details, which the parser must not extract anyway.
"""

import io
import re
import zipfile
from dataclasses import dataclass
from typing import Literal

import docx
from pypdf import PdfReader
from pypdf.errors import PdfReadError

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

MAX_PDF_PAGES = 20  # longer documents are read up to this page
MAX_CHARS = 40_000  # text beyond this is not sent to the model
MIN_MEANINGFUL_CHARS = 200  # fewer letters/digits than this: treat as scanned / empty
MAX_DOCX_UNCOMPRESSED_BYTES = 30 * 1024 * 1024  # zip-bomb guard
MAX_DOCX_ENTRIES = 2_000

ExtractError = Literal["no_text", "encrypted", "invalid_file", "too_large"]


@dataclass(frozen=True)
class Extracted:
    text: str | None
    error: ExtractError | None = None
    truncated: bool = False


def extract_text(data: bytes, content_type: str) -> Extracted:
    """Extract plain text. Never raises for bad input: problems come back as `error`."""
    try:
        if content_type == PDF:
            raw = _pdf_text(data)
        elif content_type == DOCX:
            raw = _docx_text(data)
        else:
            return Extracted(text=None, error="invalid_file")
    except _UnreadableFileError as failure:
        return Extracted(text=None, error=failure.code)
    except Exception:  # any parser crash on hostile input is just an unreadable file
        return Extracted(text=None, error="invalid_file")

    text = _normalise(raw)
    if sum(ch.isalnum() for ch in text) < MIN_MEANINGFUL_CHARS:
        return Extracted(text=None, error="no_text")
    if len(text) > MAX_CHARS:
        return Extracted(text=text[:MAX_CHARS], truncated=True)
    return Extracted(text=text)


class _UnreadableFileError(Exception):
    def __init__(self, code: ExtractError) -> None:
        super().__init__(code)
        self.code: ExtractError = code


def _pdf_text(data: bytes) -> str:
    if not data.startswith(b"%PDF-"):
        raise _UnreadableFileError("invalid_file")
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise _UnreadableFileError("encrypted")
        parts: list[str] = []
        for index, page in enumerate(reader.pages):
            if index >= MAX_PDF_PAGES:
                break
            parts.append(page.extract_text() or "")
            if sum(len(p) for p in parts) > MAX_CHARS:
                break
    except PdfReadError as exc:
        raise _UnreadableFileError("invalid_file") from exc
    return "\n".join(parts)


def _docx_text(data: bytes) -> str:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise _UnreadableFileError("invalid_file") from exc
    with archive:
        entries = archive.infolist()
        if len(entries) > MAX_DOCX_ENTRIES:
            raise _UnreadableFileError("too_large")
        if sum(entry.file_size for entry in entries) > MAX_DOCX_UNCOMPRESSED_BYTES:
            raise _UnreadableFileError("too_large")
        if "word/document.xml" not in archive.namelist():
            raise _UnreadableFileError("invalid_file")

    document = docx.Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


_SPACES = re.compile(r"[ \t\u00a0]+")
_BLANK_LINES = re.compile(r"\n{3,}")


def _normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = _SPACES.sub(" ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    return _BLANK_LINES.sub("\n\n", text).strip()
